"""Office / misc document parsing via MarkItDown (MIT).

Covers DOCX, PPTX, XLSX, HTML, CSV, ZIP and images. MarkItDown is
used for its markdown output; markdown pipe-tables in that output are also
converted into structured Page.tables.

(MD/TXT are handled by parsers.text with the stdlib, so the base install
can ingest them without this extra.)

Special cases handled here instead of plain MarkItDown:
  * XLSX -> one Page per sheet (via openpyxl, a MarkItDown extra), so sheet
    structure survives instead of one merged blob.
  * CSV  -> one Page with one table (stdlib csv module).
  * ZIP  -> members are extracted and parsed recursively (depth 1).
  * images (png/jpg/...) -> single Page with needs_ocr=True: no text layer,
    escalated to the OCR/vision tier by the router. No textless stub is
    emitted as content.

Requires the ``office`` extra: ``pip install "okfsmith[office]"``.
"""

from __future__ import annotations

from typing import TYPE_CHECKING

if TYPE_CHECKING:
    from . import Page, ParsedDocument

import csv
import logging
import tempfile
import zipfile
from pathlib import Path

from .ziputil import safe_extract

log = logging.getLogger(__name__)

_IMAGE_SUFFIXES = frozenset(
    {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tiff", ".tif", ".webp"}
)


def _markitdown():
    try:
        from markitdown import MarkItDown
    except ImportError as exc:
        raise RuntimeError(
            "markitdown is not installed; install okfsmith with the office "
            'extra (`pip install "okfsmith[office]"`) to parse this file type'
        ) from exc
    return MarkItDown()


def markdown_tables(md_text: str) -> list[list[list[str]]]:
    """Extract GitHub-style pipe tables from markdown text.

    Returns a list of tables; each table is a list of rows; each row a list
    of cell strings. The ``|---|---|`` separator row is dropped.
    """
    tables: list[list[list[str]]] = []
    current: list[list[str]] = []

    def flush() -> None:
        nonlocal current
        if current:
            tables.append(current)
            current = []

    for raw_line in md_text.splitlines():
        line = raw_line.strip()
        if line.startswith("|") and line.endswith("|"):
            cells = [c.strip() for c in line.strip("|").split("|")]
            if cells and all(set(c) <= set("-: ") and c for c in cells):
                continue  # separator row like |---|---|
            current.append(cells)
        else:
            flush()
    flush()
    return tables


def _page_from_markdown(number: int, text: str) -> Page:
    from . import Page

    text = (text or "").strip()
    return Page(
        number=number,
        text=text,
        tables=markdown_tables(text),
        needs_ocr=not bool(text),
    )


def _parse_xlsx(path: Path) -> ParsedDocument:
    """One Page per sheet, via openpyxl (no MarkItDown merge)."""
    from . import Page, ParsedDocument

    try:
        import openpyxl
    except ImportError:
        log.warning("openpyxl missing; falling back to MarkItDown for %s", path.name)
        return _parse_with_markitdown(path)

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    pages: list[Page] = []
    for i, name in enumerate(wb.sheetnames, start=1):
        ws = wb[name]
        rows: list[list[str]] = []
        for row in ws.iter_rows(values_only=True):
            cells = ["" if v is None else str(v) for v in row]
            if any(cells):
                rows.append(cells)
        if not rows:
            pages.append(Page(number=i, text="", tables=[], needs_ocr=False))
            continue
        md = "\n".join("| " + " | ".join(r) + " |" for r in rows)
        pages.append(Page(number=i, text=f"## {name}\n\n{md}", tables=[rows]))
    wb.close()
    return ParsedDocument(
        pages=pages,
        meta={"source": str(path), "tier": "TIER1_LOCAL", "parser": "openpyxl"},
    )


def _parse_csv(path: Path) -> ParsedDocument:
    from . import ParsedDocument

    with open(path, newline="", encoding="utf-8-sig") as fh:
        rows = [list(row) for row in csv.reader(fh) if any(row)]
    md = "\n".join("| " + " | ".join(r) + " |" for r in rows)
    page = _page_from_markdown(1, md)
    page.tables = [rows] if rows else []
    return ParsedDocument(
        pages=[page],
        meta={"source": str(path), "tier": "TIER1_LOCAL", "parser": "csv"},
    )


def _parse_zip(path: Path, _depth: int = 0) -> ParsedDocument:
    from . import ParsedDocument

    pages: list = []
    with tempfile.TemporaryDirectory(prefix="okfsmith-zip-") as tmp:
        with zipfile.ZipFile(path) as zf:
            safe_extract(zf, tmp)
        members = sorted(
            p for p in Path(tmp).rglob("*") if p.is_file() and not p.name.startswith(".")
        )
        for member in members:
            rel = member.relative_to(tmp).as_posix()
            try:
                if member.suffix.lower() == ".zip" or _depth >= 1:
                    continue
                sub = parse_office(member, _depth=_depth + 1)
            except Exception as exc:  # one bad member must not kill the zip
                log.warning("zip member %s skipped: %s", rel, exc)
                continue
            for pg in sub.pages:
                pg.text = f"[{rel}]\n\n{pg.text}" if pg.text else f"[{rel}]"
                pg.number = len(pages) + 1
                pages.append(pg)
    return ParsedDocument(
        pages=pages,
        meta={"source": str(path), "tier": "TIER1_LOCAL", "parser": "zip"},
    )


def _parse_image(path: Path) -> ParsedDocument:
    """Images have no text layer: flag for the OCR/vision tier, no stub text."""
    from . import Page, ParsedDocument

    return ParsedDocument(
        pages=[Page(number=1, text="", tables=[], needs_ocr=True)],
        meta={
            "source": str(path),
            "tier": "TIER3_OCR",
            "parser": "none",
            "note": "image file: no text layer; needs OCR/vision tier",
        },
    )


def _parse_with_markitdown(path: Path) -> ParsedDocument:
    from . import ParsedDocument

    md = _markitdown()
    result = md.convert(str(path))
    text = (result.text_content or "").strip()
    page = _page_from_markdown(1, text)
    return ParsedDocument(
        pages=[page],
        meta={"source": str(path), "tier": "TIER1_LOCAL", "parser": "markitdown"},
    )


def parse_office(path: str | Path, _depth: int = 0) -> ParsedDocument:
    """Parse an office/misc file into a ParsedDocument."""
    p = Path(path)
    suffix = p.suffix.lower()
    if suffix in _IMAGE_SUFFIXES:
        return _parse_image(p)
    if suffix == ".xlsx":
        return _parse_xlsx(p)
    if suffix == ".csv":
        return _parse_csv(p)
    if suffix == ".zip":
        return _parse_zip(p, _depth=_depth)
    return _parse_with_markitdown(p)


__all__ = ["parse_office", "markdown_tables"]
